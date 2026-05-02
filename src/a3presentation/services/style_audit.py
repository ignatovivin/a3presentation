from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from a3presentation.domain.presentation import PresentationPlan, SlideSpec
from a3presentation.domain.template import PlaceholderSpec, TemplateManifest


STYLE_GEOMETRY_TOLERANCE_EMU = 90000


@dataclass(frozen=True)
class StyleViolation:
    slide_index: int
    title: str
    rule: str
    details: str


def audit_presentation_styles(
    output_path: Path,
    plan: PresentationPlan,
    manifest: TemplateManifest,
) -> list[StyleViolation]:
    presentation = Presentation(str(output_path))
    violations: list[StyleViolation] = []

    for slide_index, slide_spec in enumerate(plan.slides, start=1):
        if slide_index > len(presentation.slides):
            break
        layout = _layout_for_slide(slide_spec, manifest)
        if layout is None:
            continue
        slide = presentation.slides[slide_index - 1]
        expected_background_color = _normalize_hex_color(
            getattr(getattr(layout, "background_style", None), "fill_color", None)
        )
        if expected_background_color:
            actual_background_color = _slide_background_color(slide)
            if actual_background_color != expected_background_color:
                violations.append(
                    StyleViolation(
                        slide_index=slide_index,
                        title=slide_spec.title or "",
                        rule="background_fill_color_mismatch",
                        details=f"actual={actual_background_color} expected={expected_background_color}",
                    )
                )
        violations.extend(_table_style_violations(slide, slide_spec, manifest, slide_index))
        violations.extend(_chart_style_violations(slide, slide_spec, manifest, slide_index))
        placeholders = {
            shape.placeholder_format.idx: shape
            for shape in slide.placeholders
            if getattr(shape, "is_placeholder", False)
        }
        for placeholder_spec in layout.placeholders:
            expected_text_color = _normalize_hex_color(getattr(getattr(placeholder_spec, "text_style", None), "color", None))
            shape_style = getattr(placeholder_spec, "shape_style", None)
            expected_fill_color = _normalize_hex_color(getattr(shape_style, "fill_color", None)) if shape_style is not None else None
            expected_line_color = _normalize_hex_color(getattr(shape_style, "line_color", None)) if shape_style is not None else None
            if not expected_text_color and not expected_fill_color and not expected_line_color:
                continue

            shape = _shape_for_placeholder(slide, placeholders, placeholder_spec)
            if shape is None:
                violations.append(
                    StyleViolation(
                        slide_index=slide_index,
                        title=slide_spec.title or "",
                        rule="style_target_missing",
                        details=f"placeholder={placeholder_spec.name} idx={placeholder_spec.idx}",
                    )
                )
                continue

            if expected_text_color:
                actual_text_colors = _text_run_colors(shape)
                if actual_text_colors and any(color != expected_text_color for color in actual_text_colors):
                    violations.append(
                        StyleViolation(
                            slide_index=slide_index,
                            title=slide_spec.title or "",
                            rule="text_color_mismatch",
                            details=(
                                f"placeholder={placeholder_spec.name} "
                                f"actual={sorted(actual_text_colors)} expected={expected_text_color}"
                            ),
                        )
                    )
            if expected_fill_color:
                actual_fill_color = _shape_fill_color(shape)
                if actual_fill_color != expected_fill_color:
                    violations.append(
                        StyleViolation(
                            slide_index=slide_index,
                            title=slide_spec.title or "",
                            rule="shape_fill_color_mismatch",
                            details=f"placeholder={placeholder_spec.name} actual={actual_fill_color} expected={expected_fill_color}",
                        )
                    )
            if expected_line_color:
                actual_line_color = _shape_line_color(shape)
                if actual_line_color != expected_line_color:
                    violations.append(
                        StyleViolation(
                            slide_index=slide_index,
                            title=slide_spec.title or "",
                            rule="shape_line_color_mismatch",
                            details=f"placeholder={placeholder_spec.name} actual={actual_line_color} expected={expected_line_color}",
                        )
                    )

    return violations


def _layout_for_slide(slide_spec: SlideSpec, manifest: TemplateManifest):
    key = slide_spec.preferred_layout_key
    if not key and slide_spec.render_target is not None:
        key = slide_spec.render_target.key
    if not key:
        return None
    return next((layout for layout in manifest.layouts if layout.key == key), None)


def _shape_for_placeholder(slide, placeholders: dict[int, object], placeholder_spec: PlaceholderSpec):
    if placeholder_spec.idx is not None and placeholder_spec.idx in placeholders:
        return placeholders[placeholder_spec.idx]
    if placeholder_spec.shape_name:
        named = next((shape for shape in slide.shapes if getattr(shape, "name", None) == placeholder_spec.shape_name), None)
        if named is not None:
            return named
    return _shape_matching_geometry(slide, placeholder_spec)


def _shape_matching_geometry(slide, placeholder_spec: PlaceholderSpec):
    expected = (
        placeholder_spec.left_emu,
        placeholder_spec.top_emu,
        placeholder_spec.width_emu,
        placeholder_spec.height_emu,
    )
    if None in expected:
        return None
    for shape in slide.shapes:
        actual = (
            _shape_dimension(shape, "left"),
            _shape_dimension(shape, "top"),
            _shape_dimension(shape, "width"),
            _shape_dimension(shape, "height"),
        )
        if None in actual:
            continue
        if all(abs(int(actual_value) - int(expected_value)) <= STYLE_GEOMETRY_TOLERANCE_EMU for actual_value, expected_value in zip(actual, expected)):
            return shape
    return None


def _shape_dimension(shape, attr: str) -> int | None:
    value = getattr(shape, attr, None)
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _text_run_colors(shape) -> set[str]:
    if not getattr(shape, "has_text_frame", False):
        return set()
    return _text_frame_run_colors(shape.text_frame)


def _text_frame_run_colors(text_frame) -> set[str]:
    colors: set[str] = set()
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            color = _normalize_hex_color(getattr(getattr(run.font, "color", None), "rgb", None))
            if color:
                colors.add(color)
    return colors


def _shape_fill_color(shape) -> str | None:
    try:
        return _normalize_hex_color(shape.fill.fore_color.rgb)
    except Exception:
        return None


def _shape_line_color(shape) -> str | None:
    try:
        return _normalize_hex_color(shape.line.color.rgb)
    except Exception:
        return None


def _slide_background_color(slide) -> str | None:
    try:
        return _normalize_hex_color(slide.background.fill.fore_color.rgb)
    except Exception:
        return _background_color_from_xml(getattr(slide._element.cSld, "bg", None))


def _background_color_from_xml(background) -> str | None:
    if background is None:
        return None
    for element in background.iter():
        if element.tag.endswith("}srgbClr"):
            return _normalize_hex_color(element.get("val"))
    return None


def _table_style_violations(slide, slide_spec: SlideSpec, manifest: TemplateManifest, slide_index: int) -> list[StyleViolation]:
    if getattr(slide_spec.kind, "value", slide_spec.kind) != "table":
        return []
    table_shape = next((shape for shape in slide.shapes if getattr(shape, "has_table", False)), None)
    if table_shape is None:
        return []

    expected_header_fill = _component_behavior_color(manifest, "table", "header_fill_color") or _design_token_color(
        manifest,
        "table_header_fill_color",
    )
    expected_header_text = _component_behavior_color(manifest, "table", "header_text_color") or _design_token_color(
        manifest,
        "table_header_text_color",
    )
    if not expected_header_fill and not expected_header_text:
        return []

    violations: list[StyleViolation] = []
    table = table_shape.table
    if len(table.rows) <= 0:
        return []
    for col_index, cell in enumerate(table.rows[0].cells):
        if expected_header_fill:
            actual_fill = _table_cell_fill_color(cell)
            if actual_fill != expected_header_fill:
                violations.append(
                    StyleViolation(
                        slide_index=slide_index,
                        title=slide_spec.title or "",
                        rule="table_header_fill_color_mismatch",
                        details=f"col={col_index} actual={actual_fill} expected={expected_header_fill}",
                    )
                )
        if expected_header_text:
            actual_text_colors = _table_cell_text_colors(cell)
            if actual_text_colors and any(color != expected_header_text for color in actual_text_colors):
                violations.append(
                    StyleViolation(
                        slide_index=slide_index,
                        title=slide_spec.title or "",
                        rule="table_header_text_color_mismatch",
                        details=f"col={col_index} actual={sorted(actual_text_colors)} expected={expected_header_text}",
                    )
                )
    return violations


def _component_behavior_color(manifest: TemplateManifest, component_key: str, token: str) -> str | None:
    component_style = manifest.component_styles.get(component_key)
    if component_style is None:
        return None
    return _normalize_hex_color(component_style.behavior_tokens.get(token))


def _design_token_color(manifest: TemplateManifest, token: str) -> str | None:
    return _normalize_hex_color(manifest.design_tokens.get(token))


def _table_cell_fill_color(cell) -> str | None:
    try:
        return _normalize_hex_color(cell.fill.fore_color.rgb)
    except Exception:
        return _table_cell_color_from_xml(cell, "solidFill")


def _table_cell_text_colors(cell) -> set[str]:
    return _text_frame_run_colors(cell.text_frame)


def _table_cell_color_from_xml(cell, container_name: str) -> str | None:
    for element in cell._tc.iter():
        if not element.tag.endswith(f"}}{container_name}"):
            continue
        for child in element.iter():
            if child.tag.endswith("}srgbClr"):
                return _normalize_hex_color(child.get("val"))
    return None


def _chart_style_violations(slide, slide_spec: SlideSpec, manifest: TemplateManifest, slide_index: int) -> list[StyleViolation]:
    if getattr(slide_spec.kind, "value", slide_spec.kind) != "chart" or slide_spec.chart is None:
        return []
    chart_shape = next((shape for shape in slide.shapes if getattr(shape, "has_chart", False)), None)
    if chart_shape is None:
        return []
    expected_palette = _chart_palette_colors(manifest)
    if not expected_palette:
        return []
    expected_series_colors = [
        expected_palette[index % len(expected_palette)]
        for index, _series in enumerate(slide_spec.chart.series)
    ]
    if not expected_series_colors:
        return []

    chart_xml = chart_shape.chart._chartSpace.xml.upper()
    violations: list[StyleViolation] = []
    for series_index, expected_color in enumerate(expected_series_colors):
        if expected_color not in chart_xml:
            violations.append(
                StyleViolation(
                    slide_index=slide_index,
                    title=slide_spec.title or "",
                    rule="chart_series_color_missing",
                    details=f"series={series_index} expected={expected_color}",
                )
            )
    return violations


def _chart_palette_colors(manifest: TemplateManifest) -> list[str]:
    component_colors = [
        _component_behavior_color(manifest, "chart", f"rank_color_{index}")
        for index in range(1, 5)
    ]
    if any(color is not None for color in component_colors):
        return [color for color in component_colors if color is not None]
    design_colors = [
        _design_token_color(manifest, f"chart_rank_color_{index}")
        for index in range(1, 5)
    ]
    return [color for color in design_colors if color is not None]


def _normalize_hex_color(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip().lstrip("#").upper()
    if len(text) != 6:
        return None
    if any(char not in "0123456789ABCDEF" for char in text):
        return None
    return text
